"""Extraction de texte et découpage en fragments pour l'indexation."""
import csv, io, json, zipfile

CHUNK = 1200
OVERLAP = 150
MAX_ZIP_ENTRIES = 200
MAX_ZIP_TOTAL = 50 * 1024 * 1024  # protège des zip bombs


def chunk_text(text: str, size: int = CHUNK, overlap: int = OVERLAP) -> list[str]:
    text = text.strip()
    if not text:
        return []
    out, i = [], 0
    while i < len(text):
        out.append(text[i:i + size])
        if i + size >= len(text):
            break
        i += size - overlap
    return out


def extract_text(ext: str, data: bytes) -> str:
    """Retourne le texte indexable ; '' si le type n'a pas de texte (images)."""
    if ext in ("png", "jpg", "jpeg", "webp"):
        return ""
    if ext == "json":
        return json.dumps(json.loads(data.decode()), ensure_ascii=False, indent=1)
    if ext == "csv":
        rows = csv.reader(io.StringIO(data.decode("utf-8", "replace")))
        return "\n".join(" | ".join(r) for _, r in zip(range(5000), rows))
    if ext == "pdf":
        from pypdf import PdfReader  # import tardif
        return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
    if ext == "docx":
        from docx import Document
        return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
    if ext == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"# {ws.title}")
            for _, row in zip(range(5000), ws.iter_rows(values_only=True)):
                lines.append(" | ".join("" if c is None else str(c) for c in row))
        return "\n".join(lines)
    if ext == "zip":  # on liste le contenu, sans jamais extraire sur disque
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            infos = z.infolist()
            if len(infos) > MAX_ZIP_ENTRIES or sum(i.file_size for i in infos) > MAX_ZIP_TOTAL:
                raise ValueError("Archive trop volumineuse")
            return "\n".join(f"{i.filename} ({i.file_size} o)" for i in infos)
    return data.decode("utf-8", "replace")
